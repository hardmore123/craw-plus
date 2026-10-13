"""六个日本电视品牌 SPEC 的周度刷新编排器。

默认执行：spec-crawl → spec-export → xlsx 结构/中文覆盖校验。
Windows 和 Linux 的定时任务都调用本文件，避免两套刷新逻辑长期漂移。
"""
from __future__ import annotations

import argparse
import os
import subprocess
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
BRANDS: tuple[tuple[str, str], ...] = (
    ("hisense_jp", "SPEC_海信日本.xlsx"),
    ("sony_jp", "SPEC_索尼日本.xlsx"),
    ("regza_jp", "SPEC_REGZA日本.xlsx"),
    ("panasonic_jp", "SPEC_松下日本.xlsx"),
    ("tcl_jp", "SPEC_TCL日本.xlsx"),
    ("sharp_jp", "SPEC_夏普日本.xlsx"),
)
BRAND_CODES = {code for code, _ in BRANDS}


def _region_brands(region_code: str) -> tuple[tuple[str, str], ...]:
    """从集中区域配置构造 (品牌 code, 输出 Excel 文件名) 列表（用于 --region）。

    非日本线的英文/西语规格无法满足日本模板的中文覆盖校验，因此文件名带区域后缀，
    且 main() 对非 jp 区域跳过中文覆盖校验。"""
    try:
        if str(ROOT) not in sys.path:
            sys.path.insert(0, str(ROOT))
        from overseas.regions_config import get_region
    except Exception:
        return ()
    region = get_region(region_code)
    if not region:
        return ()
    out: list[tuple[str, str]] = []
    for b in region.brands:
        out.append((b.code, f"SPEC_{b.name}_{region.code.upper()}.xlsx"))
    return tuple(out)


def _path_arg(value: str | None, default: Path) -> Path:
    if not value:
        return default.resolve()
    p = Path(value).expanduser()
    return (ROOT / p).resolve() if not p.is_absolute() else p.resolve()


def _run(command: list[str], env: dict[str, str], timeout: float | None = None) -> int:
    print("\n[RUN]", " ".join(command), flush=True)
    if timeout is None:
        try:
            timeout = max(1.0, float(env.get("OVERSEAS_SPEC_TIMEOUT_SECONDS", "1800")))
        except (TypeError, ValueError):
            timeout = 1800.0
    try:
        completed = subprocess.run(command, cwd=ROOT, env=env, timeout=timeout)
    except subprocess.TimeoutExpired:
        print(f"[TIMEOUT] 子进程超过 {timeout:.0f}s，已终止: {command[-1]}", flush=True)
        return 124
    return completed.returncode


def _validate_xlsx(path: Path, require_zh: bool = True) -> bool:
    """校验交付 Excel 的基本结构和（可选）中文覆盖率。

    require_zh=False 用于英文/西语规格线（如加拿大），只校验表头与型号/规格行存在，
    不因项目未翻译中文而判失败。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("[FAIL] 缺少 openpyxl，无法校验 xlsx", flush=True)
        return False
    if not path.is_file() or path.stat().st_size == 0:
        print(f"[FAIL] 输出文件不存在或为空: {path}", flush=True)
        return False
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as exc:
        print(f"[FAIL] 无法读取 {path.name}: {type(exc).__name__}: {exc}", flush=True)
        return False
    finally:
        try:
            wb.close()
        except Exception:
            pass

    if not rows or tuple(rows[0][:3]) != ("区分", "项目", "项目（中文）"):
        print(f"[FAIL] {path.name} 表头不符合 SPEC 格式", flush=True)
        return False
    data = [r for r in rows[1:] if len(r) > 1 and str(r[1] or "").strip()]
    models = [str(x).strip() for x in rows[0][3:] if str(x or "").strip()]
    missing = [str(r[1]) for r in data if not str(r[2] or "").strip()]
    coverage = (len(data) - len(missing)) / max(len(data), 1) * 100
    print(f"[XLSX] {path.name}: rows={len(data)} models={len(models)} "
          f"中文={coverage:.1f}%", flush=True)
    if not models or not data:
        print(f"[FAIL] {path.name} 没有型号列或规格行", flush=True)
        return False
    if require_zh and missing:
        print(f"[FAIL] {path.name} 有 {len(missing)} 个未翻译项目: {missing[:10]}", flush=True)
        return False
    return True


def build_parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(description="TV SPEC 周度刷新（默认日本六品牌，--region 切换其它线）")
    ap.add_argument("--site", action="append",
                    help="只刷新指定品牌 code，可重复；缺省刷新该区域全部品牌")
    ap.add_argument("--region", default="jp",
                    help="区域线：jp=日本六品牌（默认），ca=加拿大六品牌等")
    ap.add_argument("--db", help="SQLite 路径，缺省使用 OVERSEAS_DB_PATH 或 data/overseas.db")
    ap.add_argument("--output-dir", help="Excel 输出目录，缺省为项目根目录")
    ap.add_argument("--export-only", action="store_true",
                    help="跳过采集，只从数据库重新导出并校验 Excel")
    ap.add_argument("--allow-incomplete", action="store_true",
                    help="允许 Excel 存在未翻译项目（默认将其视为失败）")
    ap.add_argument("--fail-fast", action="store_true",
                    help="首个品牌失败后立即停止；默认继续执行其余品牌")
    return ap


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    region = (args.region or "jp").strip().lower()
    if region == "jp":
        all_brands = BRANDS
    else:
        all_brands = _region_brands(region)
        if not all_brands:
            print(f"[FAIL] 未知或空的区域: {region}", flush=True)
            return 1
    all_codes = {code for code, _ in all_brands}
    selected = set(args.site) if args.site else all_codes
    unknown = selected - all_codes
    if unknown:
        print(f"[FAIL] 区域 {region} 不包含品牌: {sorted(unknown)}", flush=True)
        return 1
    brands = [(code, filename) for code, filename in all_brands if code in selected]
    # 非日本线规格是英文/西语，跳过中文覆盖校验（仍校验表头与型号/规格行存在）。
    skip_zh = region != "jp"
    db = _path_arg(args.db, Path(os.environ.get("OVERSEAS_DB_PATH") or ROOT / "data" / "overseas.db"))
    output_dir = _path_arg(args.output_dir, ROOT)
    output_dir.mkdir(parents=True, exist_ok=True)

    env = os.environ.copy()
    # REGZA/TCL 页面较重；用户显式设置时尊重其配置。
    env.setdefault("OVERSEAS_NAV_TIMEOUT", "60")
    env.setdefault("PYTHONIOENCODING", "utf-8")
    failures: list[str] = []

    print(f"[SPEC] 品牌数={len(brands)} db={db} output={output_dir}", flush=True)
    for code, filename in brands:
        out = output_dir / filename
        if not args.export_only:
            rc = _run([sys.executable, "-m", "overseas.cli", "--db", str(db),
                       "spec-crawl", "--site", code], env)
            if rc != 0:
                failures.append(f"{code}: spec-crawl rc={rc}")
                if args.fail_fast:
                    break
                # 即使本轮部分失败，也继续导出当前数据库中可用/保留的快照；
                # 最终仍由 failures 返回非零，避免旧 Excel 被静默当成完整成功。

        rc = _run([sys.executable, "-m", "overseas.cli", "--db", str(db),
                   "spec-export", "--site", code, "--format", "xlsx",
                   "--out", str(out)], env)
        if rc != 0:
            failures.append(f"{code}: spec-export rc={rc}")
            if args.fail_fast:
                break
            continue
        if not _validate_xlsx(out, require_zh=not skip_zh) and not args.allow_incomplete:
            failures.append(f"{code}: xlsx validation failed")
            if args.fail_fast:
                break

    # 即使某个品牌采集失败，也导出当前运行可见的系列状态，便于定时任务直接审计。
    status_out = output_dir / f"SPEC_系列状态_{datetime.now():%Y%m%d}"
    status_command = [
        sys.executable,
        str(ROOT / "scripts" / "export_spec_status.py"),
        "--db", str(db),
        "--format", "all",
        "--out", str(status_out),
    ]
    for code, _ in brands:
        status_command.extend(["--site", code])
    rc = _run(status_command, env)
    if rc != 0:
        failures.append(f"spec-status-export rc={rc}")

    if failures:
        print("\n[FAIL] SPEC 刷新未完整成功:", flush=True)
        for item in failures:
            print(" -", item, flush=True)
        return 1
    print("\n[OK] SPEC 刷新、导出、状态汇总和校验全部完成", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
