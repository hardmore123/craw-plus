"""汇总日本六品牌 SPEC 阶段发现的电视型号，供后续平台搜索评价和价格使用。

默认从 spec_series.models 读取当前数据库中的型号列表，一行输出一个“品牌 + 标准化型号”。
输出同时保留官网原始型号、系列、来源 URL、最新 SPEC 周和重复审计信息。

示例：
    python scripts/export_spec_models.py
    python scripts/export_spec_models.py --format all
    python scripts/export_spec_models.py --site sharp_jp --format xlsx --out data/sharp_models.xlsx
    python scripts/export_spec_models.py --include-duplicates --format csv
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sqlite3
import sys
import unicodedata
from collections import Counter, OrderedDict, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))
from overseas.spec_identity import display_series_name

DEFAULT_DB = PROJECT_ROOT / "data" / "overseas.db"

BRANDS = (
    "hisense_jp",
    "sony_jp",
    "regza_jp",
    "panasonic_jp",
    "tcl_jp",
    "sharp_jp",
)

BRAND_SEARCH_NAME = {
    "hisense_jp": "Hisense",
    "sony_jp": "Sony BRAVIA",
    "regza_jp": "REGZA",
    "panasonic_jp": "Panasonic VIERA",
    "tcl_jp": "TCL",
    "sharp_jp": "SHARP AQUOS",
}


def _region_brand_codes(region_code: str) -> list[str]:
    """从集中区域配置取某条线的品牌 code 列表（用于 --region）。"""
    try:
        sys.path.insert(0, str(PROJECT_ROOT))
        from overseas.regions_config import get_region
    except Exception:
        return []
    region = get_region(region_code)
    return region.brand_codes() if region else []


def _region_search_names(region_code: str) -> dict[str, str]:
    """区域品牌的搜索展示名（用于 search_query 列）。"""
    try:
        sys.path.insert(0, str(PROJECT_ROOT))
        from overseas.regions_config import get_region
    except Exception:
        return {}
    region = get_region(region_code)
    if not region:
        return {}
    return {b.code: b.query_name for b in region.brands}

BASE_COLUMNS = [
    "brand_code",
    "brand_name",
    "model_raw",
    "model_normalized",
    "search_query",
    "series",
    "series_count",
    "source_url",
    "source_scope",
    "latest_spec_week",
    "series_first_seen",
    "series_last_seen",
    "observed_count",
    "duplicate_in_series",
    "duplicate_occurrences",
    "cross_brand_count",
    "cross_brand_same_normalized",
]
DUPLICATE_COLUMNS = ["occurrence_index_in_series", "occurrence_count_in_series"]


def normalize_model(value: str) -> str:
    """生成用于去重和平台搜索匹配的标准化型号，不替代原始型号。"""
    text = unicodedata.normalize("NFKC", value).upper()
    return re.sub(r"[^A-Z0-9]+", "", text)


def clean_model(value: Any) -> str:
    if not isinstance(value, str):
        return ""
    return " ".join(value.strip().split())


def source_scope(url: str) -> str:
    """说明 spec_series.url 是系列页还是多级流程中的第一个型号来源。"""
    lower = url.lower()
    if "product_spec." in lower:
        return "model_json_first"
    if "/viera/products/tv-" in lower:
        return "first_model_spec"
    return "series_spec"


def parse_models(raw: Any, *, brand: str, series: str) -> list[str]:
    if raw is None or raw == "":
        print(f"[WARN] {brand}/{series}: models 为空", file=sys.stderr)
        return []
    try:
        values = json.loads(raw) if isinstance(raw, str) else raw
    except (TypeError, json.JSONDecodeError) as exc:
        print(f"[WARN] {brand}/{series}: models JSON 无法解析: {exc}", file=sys.stderr)
        return []
    if not isinstance(values, list):
        print(f"[WARN] {brand}/{series}: models 不是数组，已跳过", file=sys.stderr)
        return []
    result = [clean_model(value) for value in values]
    result = [value for value in result if value]
    if not result:
        print(f"[WARN] {brand}/{series}: models 数组没有有效型号", file=sys.stderr)
    return result


def _min_nonempty(current: str, value: str) -> str:
    if not value:
        return current
    if not current or value < current:
        return value
    return current


def _max_nonempty(current: str, value: str) -> str:
    if not value:
        return current
    if not current or value > current:
        return value
    return current


def load_series_rows(
    conn: sqlite3.Connection,
    sites: Iterable[str],
) -> tuple[list[dict[str, Any]], int]:
    sites = tuple(sites)
    placeholders = ",".join("?" for _ in sites)
    query = f"""
        SELECT
            s.id,
            s.brand,
            s.brand_name,
            s.series,
            s.url,
            s.models,
            s.first_seen,
            s.last_seen,
            MAX(r.captured_week) AS latest_spec_week
        FROM spec_series AS s
        LEFT JOIN spec_row AS r ON r.series_id = s.id
        WHERE s.brand IN ({placeholders})
        GROUP BY s.id, s.brand, s.brand_name, s.series, s.url,
                 s.models, s.first_seen, s.last_seen
        ORDER BY s.brand, s.id
    """
    rows = conn.execute(query, sites).fetchall()
    occurrences: list[dict[str, Any]] = []
    for row in rows:
        brand = str(row["brand"] or "")
        series = str(row["series"] or "")
        raw_models = parse_models(row["models"], brand=brand, series=series)
        counts = Counter(normalize_model(model) for model in raw_models)
        positions = defaultdict(int)
        for position, model_raw in enumerate(raw_models, start=1):
            model_normalized = normalize_model(model_raw)
            if not model_normalized:
                print(f"[WARN] {brand}/{series}: 型号 {model_raw!r} 标准化后为空，已跳过", file=sys.stderr)
                continue
            positions[model_normalized] += 1
            occurrences.append(
                {
                    "brand_code": brand,
                    "brand_name": str(row["brand_name"] or BRAND_SEARCH_NAME.get(brand, brand)),
                    "model_raw": model_raw,
                    "model_normalized": model_normalized,
                    "series": display_series_name(series),
                    "source_url": str(row["url"] or ""),
                    "source_scope": source_scope(str(row["url"] or "")),
                    "latest_spec_week": str(row["latest_spec_week"] or ""),
                    "series_first_seen": str(row["first_seen"] or ""),
                    "series_last_seen": str(row["last_seen"] or ""),
                    "occurrence_index_in_series": positions[model_normalized],
                    "occurrence_count_in_series": counts[model_normalized],
                    "model_position_in_series": position,
                }
            )
    return occurrences, len(rows)


def aggregate_occurrences(occurrences: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """按品牌+标准化型号聚合，同时保留所有来源和重复信息。"""
    grouped: OrderedDict[tuple[str, str], dict[str, Any]] = OrderedDict()
    for occurrence in occurrences:
        key = (occurrence["brand_code"], occurrence["model_normalized"])
        item = grouped.setdefault(
            key,
            {
                "brand_code": occurrence["brand_code"],
                "brand_name": occurrence["brand_name"],
                "raw_models": [],
                "series": [],
                "source_urls": [],
                "source_scopes": [],
                "latest_weeks": [],
                "series_first_seen": "",
                "series_last_seen": "",
                "occurrences": [],
                "series_counts": Counter(),
            },
        )
        for field, list_key in (
            ("model_raw", "raw_models"),
            ("series", "series"),
            ("source_url", "source_urls"),
            ("source_scope", "source_scopes"),
            ("latest_spec_week", "latest_weeks"),
        ):
            value = occurrence[field]
            if value and value not in item[list_key]:
                item[list_key].append(value)
        item["series_first_seen"] = _min_nonempty(item["series_first_seen"], occurrence["series_first_seen"])
        item["series_last_seen"] = _max_nonempty(item["series_last_seen"], occurrence["series_last_seen"])
        item["occurrences"].append(occurrence)
        item["series_counts"][occurrence["series"]] += 1

    brands_by_normalized: dict[str, set[str]] = defaultdict(set)
    for brand_code, model_normalized in grouped:
        brands_by_normalized[model_normalized].add(brand_code)

    result: list[dict[str, Any]] = []
    for (brand_code, model_normalized), item in grouped.items():
        duplicate_occurrences = sum(max(0, count - 1) for count in item["series_counts"].values())
        raw_models = item["raw_models"]
        brand_search_name = BRAND_SEARCH_NAME.get(brand_code, item["brand_name"])
        cross_brand_count = len(brands_by_normalized[model_normalized])
        result.append(
            {
                "brand_code": brand_code,
                "brand_name": item["brand_name"],
                "model_raw": " | ".join(raw_models),
                "model_normalized": model_normalized,
                "search_query": f"{brand_search_name} {raw_models[0]}" if raw_models else model_normalized,
                "series": " | ".join(item["series"]),
                "series_count": len(item["series"]),
                "source_url": " | ".join(item["source_urls"]),
                "source_scope": " | ".join(item["source_scopes"]),
                "latest_spec_week": " | ".join(sorted(item["latest_weeks"])),
                "series_first_seen": item["series_first_seen"],
                "series_last_seen": item["series_last_seen"],
                "observed_count": len(item["occurrences"]),
                "duplicate_in_series": "yes" if duplicate_occurrences else "no",
                "duplicate_occurrences": duplicate_occurrences,
                "cross_brand_count": cross_brand_count,
                "cross_brand_same_normalized": "yes" if cross_brand_count > 1 else "no",
                "_occurrences": item["occurrences"],
            }
        )
    return result


def duplicate_rows(
    aggregates: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for aggregate in aggregates:
        for occurrence in aggregate["_occurrences"]:
            row = {column: aggregate[column] for column in BASE_COLUMNS}
            row["model_raw"] = occurrence["model_raw"]
            row["series"] = occurrence["series"]
            row["series_count"] = 1
            row["source_url"] = occurrence["source_url"]
            row["source_scope"] = occurrence["source_scope"]
            row["latest_spec_week"] = occurrence["latest_spec_week"]
            row["series_first_seen"] = occurrence["series_first_seen"]
            row["series_last_seen"] = occurrence["series_last_seen"]
            row["observed_count"] = occurrence["occurrence_count_in_series"]
            row["occurrence_index_in_series"] = occurrence["occurrence_index_in_series"]
            row["occurrence_count_in_series"] = occurrence["occurrence_count_in_series"]
            result.append(row)
    return result


def public_rows(aggregates: list[dict[str, Any]], include_duplicates: bool) -> list[dict[str, Any]]:
    if include_duplicates:
        return duplicate_rows(aggregates)
    return [{column: aggregate[column] for column in BASE_COLUMNS} for aggregate in aggregates]


def output_paths(out: str | None, output_format: str,
                 region: str | None = None) -> dict[str, Path]:
    extension = {"csv": ".csv", "xlsx": ".xlsx", "json": ".json", "jsonl": ".jsonl"}
    region_code = (region or "").strip().lower()
    if out:
        base = Path(out)
    elif region_code:
        base = PROJECT_ROOT / "data" / f"SPEC_型号汇总_{region_code}.csv"
    else:
        base = PROJECT_ROOT / "data" / f"SPEC_型号汇总_{datetime.now():%Y%m%d}.csv"
    if output_format == "all":
        if base.suffix.lower() in extension.values():
            base = base.with_suffix("")
        return {name: base.with_suffix(suffix) for name, suffix in extension.items() if name != "jsonl"}
    target_suffix = extension[output_format]
    if base.suffix.lower() != target_suffix:
        base = base.with_suffix(target_suffix) if base.suffix else Path(str(base) + target_suffix)
    return {output_format: base}


def ensure_parent(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def write_csv(path: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
    ensure_parent(path)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_json(path: Path, rows: list[dict[str, Any]]) -> None:
    ensure_parent(path)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(rows, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    ensure_parent(path)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def write_xlsx(path: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("XLSX 输出需要安装 openpyxl，例如: python -m pip install openpyxl==3.1.5") from exc

    ensure_parent(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SPEC型号汇总"
    sheet.freeze_panes = "A2"
    sheet.append(columns)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    for row in rows:
        sheet.append([row.get(column, "") for column in columns])
    if rows:
        sheet.auto_filter.ref = sheet.dimensions
    widths = {
        "brand_code": 16,
        "brand_name": 20,
        "model_raw": 24,
        "model_normalized": 20,
        "search_query": 30,
        "series": 22,
        "source_url": 58,
        "source_scope": 22,
    }
    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[chr(64 + index) if index <= 26 else "A"].width = widths.get(column, 18)
    workbook.save(path)


def write_outputs(paths: dict[str, Path], rows: list[dict[str, Any]], include_duplicates: bool) -> None:
    columns = BASE_COLUMNS + (DUPLICATE_COLUMNS if include_duplicates else [])
    for output_format, path in paths.items():
        if output_format == "csv":
            write_csv(path, rows, columns)
        elif output_format == "xlsx":
            write_xlsx(path, rows, columns)
        elif output_format == "json":
            write_json(path, rows)
        elif output_format == "jsonl":
            write_jsonl(path, rows)
        else:
            raise ValueError(f"不支持的输出格式: {output_format}")
        print(f"[OUT] {output_format}: {path} rows={len(rows)}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="汇总六品牌 SPEC 型号，生成平台搜索用清单")
    parser.add_argument("--db", default=str(DEFAULT_DB), help=f"SQLite 数据库路径，默认 {DEFAULT_DB}")
    parser.add_argument("--site", action="append",
                        help="只输出指定品牌 code，可重复；默认日本六品牌或 --region 指定线的全部品牌")
    parser.add_argument("--region",
                        help="按区域线取品牌（如 ca=加拿大六品牌）；与 --site 二选一，--site 优先")
    parser.add_argument("--format", choices=("csv", "xlsx", "json", "jsonl", "all"), default="csv")
    parser.add_argument("--out", help="输出文件路径；--format all 时作为文件名主干")
    parser.add_argument(
        "--include-duplicates",
        action="store_true",
        help="保留同一品牌/型号在同一系列中的重复槽位；默认按品牌+标准化型号汇总一行",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    if args.site:
        sites = tuple(args.site)
    elif args.region:
        region_brands = _region_brand_codes(args.region)
        if not region_brands:
            print(f"[ERROR] 未知或空的区域: {args.region}", file=sys.stderr)
            return 2
        sites = tuple(region_brands)
        # 合入该区域品牌的搜索展示名，供 search_query 列使用
        BRAND_SEARCH_NAME.update(_region_search_names(args.region))
    else:
        sites = tuple(BRANDS)
    db_path = Path(args.db)
    if not db_path.exists():
        print(f"[ERROR] 数据库不存在: {db_path}", file=sys.stderr)
        return 2

    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        occurrences, series_count = load_series_rows(conn, sites)
    finally:
        conn.close()

    aggregates = aggregate_occurrences(occurrences)
    rows = public_rows(aggregates, args.include_duplicates)
    raw_count = len(occurrences)
    unique_count = len(aggregates)
    duplicate_occurrences = sum(int(item["duplicate_occurrences"]) for item in aggregates)
    cross_brand_count = sum(1 for item in aggregates if item["cross_brand_same_normalized"] == "yes")
    print(
        f"[SPEC-MODELS] brands={len(sites)} series={series_count} "
        f"raw_slots={raw_count} unique_models={unique_count} "
        f"dedup_removed={raw_count - unique_count} "
        f"same_series_duplicates={duplicate_occurrences} "
        f"cross_brand_normalized={cross_brand_count}"
    )
    for brand in sites:
        brand_aggregates = [item for item in aggregates if item["brand_code"] == brand]
        brand_raw = sum(int(item["observed_count"]) for item in brand_aggregates)
        print(
            f"[BRAND] {brand}: series={len({series for item in brand_aggregates for series in item['series'].split(' | ') if series})} "
            f"raw_slots={brand_raw} unique_models={len(brand_aggregates)}"
        )

    paths = output_paths(args.out, args.format, args.region)
    write_outputs(paths, rows, args.include_duplicates)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
