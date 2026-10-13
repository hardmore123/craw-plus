"""校验导出的 SPEC Excel：表头、中文覆盖率、样例行。"""
import sys
from openpyxl import load_workbook

path = sys.argv[1] if len(sys.argv) > 1 else "SPEC_索尼日本.xlsx"
wb = load_workbook(path, read_only=True)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
header = rows[0]
data = rows[1:]
print("文件:", path, "| sheet:", ws.title)
print("总列数:", len(header), "| 数据行:", len(data))
print("表头前3列:", header[:3])
print("机型列(第4列起):", header[3:])
# 中文覆盖率：第3列(索引2)非空 / 有日文项目(第2列)的行
have_item = [r for r in data if (r[1] or "").strip()]
zh_ok = [r for r in have_item if (r[2] or "").strip()]
print(f"中文覆盖: {len(zh_ok)}/{len(have_item)} = {len(zh_ok)/max(len(have_item),1)*100:.1f}%")
# 未翻译样例
missing = [r[1] for r in have_item if not (r[2] or "").strip()]
print("未翻译项目(前20):", missing[:20])
print("\n--- 样例前15行 (区分|项目|中文|机型1值) ---")
for r in data[:15]:
    m1 = r[3] if len(r) > 3 else ""
    print(f"  {str(r[0])[:8]:8} | {str(r[1])[:22]:22} | {str(r[2])[:16]:16} | {str(m1)[:24]}")
