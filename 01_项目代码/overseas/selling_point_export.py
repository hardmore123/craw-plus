"""卖点导出：一个品牌一个 Excel，每条卖点一行。

列布局：品牌 | 系列 | 机型 | 类型 | 标题 | 描述 | 图片URL | 视频URL | LLM精炼摘要
行：该品牌所有系列的卖点条目，按系列→机型→顺序排列。

与 spec_export.py 对齐：橙色表头、冻结窗格、openpyxl 驱动。
"""
from __future__ import annotations

import json
import os

from .db import Database

HEADER = ["品牌", "系列", "机型", "类型", "标题", "描述", "图片URL", "视频URL", "LLM精炼摘要"]


def _collect(db: Database, brand: str) -> list[dict]:
    """从数据库收集卖点，返回按系列→机型→顺序排列的字典列表。"""
    rows = db.selling_point_list(brand=brand)
    out: list[dict] = []
    for r in rows:
        try:
            imgs = json.loads(r["image_urls"] or "[]")
        except (json.JSONDecodeError, TypeError):
            imgs = []
        try:
            vids = json.loads(r["video_urls"] or "[]")
        except (json.JSONDecodeError, TypeError):
            vids = []
        out.append({
            "brand": r["brand"],
            "brand_name": r["brand_name"] or r["brand"],
            "series": r["series"],
            "model": r["model"] or "",
            "point_type": r["point_type"],
            "title": r["title"] or "",
            "description": r["description"] or "",
            "image_urls": " | ".join(imgs) if imgs else "",
            "video_urls": " | ".join(vids) if vids else "",
            "refined_text": r["refined_text"] or "",
            "order": r["point_order"] or 0,
        })
    out.sort(key=lambda x: (x["series"], x["model"], x["order"]))
    return out


def export_xlsx(db: Database, brand: str, path: str) -> int:
    """导出卖点 Excel。返回导出行数。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as e:
        raise RuntimeError("导出 xlsx 需要 openpyxl：pip install openpyxl") from e

    rows = _collect(db, brand)
    brand_name = brand
    if rows:
        brand_name = rows[0]["brand_name"]

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = f"{brand_name}卖点"[:28]
    ws.append(HEADER)

    fill = PatternFill("solid", fgColor="E38B29")     # 橙色表头（对齐 spec_export）
    font = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for c in ws[1]:
        c.fill = fill
        c.font = font
        c.alignment = center

    for r in rows:
        ws.append([
            r["brand_name"],
            r["series"],
            r["model"],
            r["point_type"],
            r["title"],
            r["description"],
            r["image_urls"],
            r["video_urls"],
            r["refined_text"],
        ])

    # 列宽
    widths = [14, 14, 14, 10, 30, 50, 40, 40, 50]
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.freeze_panes = "A2"

    # 自动换行
    wrap = Alignment(vertical="top", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap

    wb.save(path)
    return len(rows)


def export_csv(db: Database, brand: str, path: str) -> int:
    """导出卖点 CSV。返回导出行数。"""
    import csv

    rows = _collect(db, brand)
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(HEADER)
        for r in rows:
            w.writerow([
                r["brand_name"], r["series"], r["model"], r["point_type"],
                r["title"], r["description"], r["image_urls"],
                r["video_urls"], r["refined_text"],
            ])
    return len(rows)
